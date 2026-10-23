"""差异计算：按主键匹配本地行和在线行（两种表共用）。"""
import csv
import io
import os
from datetime import date, datetime, time
from dataclasses import dataclass, field

from .http import LarkError


@dataclass(frozen=True)
class Change:
    key: tuple
    column: str
    old: str
    new: str
    ref: object = field(default=None, repr=False)   # 在线行的 ref，写入时定位用


class Plan:
    """diff 的结果。apply 只写 changed 和 added，写之前可以随意删减。

    默认只填空：在线已有不同值的格子放进 conflicts，不写；人工确认后
    移入 changed 才会写（全部覆盖：plan.changed += plan.conflicts）。
    """

    def __init__(self, table, key, changed, conflicts, added, unmatched):
        self._table = table
        self.key = key
        self.changed = changed      # list[Change]：在线为空、本地有值
        self.conflicts = conflicts  # list[Change]：在线已有不同的值（警告）
        self.added = added          # list[dict]：本地有、在线没有的行
        self.unmatched = unmatched  # int：在线有、本地没有的行数（只提示，不删）

    def __repr__(self):
        return (f"补空 {len(self.changed)} 格 / 新增 {len(self.added)} 行 / "
                f"不一致 {len(self.conflicts)} 格（未写入）/ 未匹配 {self.unmatched} 行")

    def apply(self):
        """写入飞书，返回写入记录 id（用于 Table.undo）。"""
        return self._table._apply(self)


def load(source):
    """CSV / xlsx 的路径或字节、或 list[dict] → (列名, 行)；值统一成去掉首尾空白的字符串。

    xlsx 读第一个（打开时显示的）工作表，第 1 行是表头。
    """
    if isinstance(source, (str, os.PathLike)):
        with open(source, "rb") as f:
            source = f.read()
    if isinstance(source, bytes) and source[:4] == b"PK\x03\x04":   # xlsx 本质是 zip
        raw, rows = _xlsx(source)
    elif isinstance(source, bytes):
        reader = csv.DictReader(io.StringIO(decode(source), newline=""))
        raw = reader.fieldnames or []
        rows = list(reader)
    else:
        rows = list(source)
        raw = list(dict.fromkeys(k for r in rows for k in r))
    header = [str(h).strip() for h in raw if h is not None and str(h).strip()]
    dupes = sorted({h for h in header if header.count(h) > 1})
    if dupes:
        raise LarkError(f"本地数据的列名重复：{'、'.join(dupes)}")
    rows = [{str(k).strip(): "" if v is None else str(v).strip()
             for k, v in r.items() if k is not None and str(k).strip()} for r in rows]
    return header, rows


def _xlsx(data):
    from openpyxl import load_workbook   # 只有读 xlsx 时才需要

    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    grid = [[_xlsx_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    if not grid:
        return [], []
    header = grid[0]
    return header, [dict(zip(header, r)) for r in grid[1:] if any(r)]


def _xlsx_cell(v):
    """按单元格显示的样子转文本：整数不带 .0，日期不带 00:00:00。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d" if v.time() == time() else "%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def decode(data):
    """UTF-8（含 BOM）优先，不行再按 GBK：Excel 中文版另存的 CSV 默认是 GBK。"""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            return data.decode("gb18030")
        except UnicodeDecodeError:
            raise LarkError("CSV 编码认不出来：请另存为 UTF-8 编码的 CSV") from None


def compute(online_rows, online_refs, local_rows, key, columns):
    """返回 (补空, 不一致, 新增的本地行, 匹配上的行数)。本地为空的格子不参与（不会清空在线值）。"""
    online = keyed(online_rows, key, online_refs)
    local = keyed(local_rows, key, local=True)
    changed, conflicts = [], []
    for k, (_, r) in local.items():
        if k not in online:
            continue
        ref, o = online[k]
        for c in columns:
            new, old = r.get(c, ""), o.get(c, "")
            if new and new != old:
                (conflicts if old else changed).append(Change(k, c, old, new, ref))
    added = [r for k, (_, r) in local.items() if k not in online]
    return changed, conflicts, added, sum(1 for k in local if k in online)


def keyed(rows, key, refs=None, local=False):
    """按主键建索引 {key: (ref, row)}。主键为空：在线行跳过，本地行报错；主键重复报错。"""
    idx, dupes, empty = {}, [], []
    for i, r in enumerate(rows):
        k = tuple(r.get(c, "") for c in key)
        if not any(k):
            empty.append(i + 1)
            continue
        if k in idx:
            dupes.append(k)
            continue
        idx[k] = (refs[i] if refs else i, r)
    where = "本地数据" if local else "在线表"
    if local and empty:
        raise LarkError(f"{where}第 {brief(map(str, empty), sep='、')} 条的主键为空")
    if dupes:
        uniq = list(dict.fromkeys(dupes))
        raise LarkError(f"{where}按「{' + '.join(key)}」匹配时有 {len(uniq)} 个值重复，例如："
                        f"{'、'.join(show_key(k) for k in uniq[:5])}。主键要能唯一确定一行，"
                        f"可以再多勾选一个字段组成联合主键")
    return idx


def show_key(k):
    return "/".join(k)


def brief(lines, n=20, sep="\n"):
    lines = list(lines)
    more = f"{sep}…共 {len(lines)} 处" if len(lines) > n else ""
    return sep.join(lines[:n]) + more
